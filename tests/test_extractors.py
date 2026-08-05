"""Tests de los extractores de PDF y DOCX.

Todos operan sobre archivos reales generados en tmp_path: un extractor probado
contra un mock de la librería de parseo no prueba nada útil.
"""

import pytest

from core.chunker import chunk_document
from core.config import ChunkingConfig
from core.extractors import EXTRACTORS, ExtractionError, extract_docx, extract_pdf
from core.models import Document

# --- PDF ---------------------------------------------------------------------


def test_pdf_text_is_extracted(tmp_path, make_pdf):
    make_pdf(tmp_path / "doc.pdf", ["Los drones requieren certificacion."])

    text = extract_pdf(tmp_path / "doc.pdf")

    assert "certificacion" in text


def test_pdf_pages_become_headers(tmp_path, make_pdf):
    """Cada página es una sección: es la referencia que un lector puede verificar."""
    make_pdf(tmp_path / "doc.pdf", ["Primera parte.", "Segunda parte."])

    text = extract_pdf(tmp_path / "doc.pdf")

    assert "## Página 1" in text
    assert "## Página 2" in text
    assert text.index("## Página 1") < text.index("## Página 2")


def test_pdf_page_headers_survive_into_citations(tmp_path, make_pdf):
    """La página extraída llega hasta la cita, que es el punto de todo esto."""
    long_text = "El reglamento de vuelo exige registro previo de la aeronave. " * 4
    make_pdf(tmp_path / "manual.pdf", [long_text])

    document = Document(source_file="manual.pdf", content=extract_pdf(tmp_path / "manual.pdf"))
    chunks = chunk_document(document, ChunkingConfig())

    assert chunks
    assert chunks[0].citation == "manual.pdf § Página 1"


def test_pdf_without_text_layer_raises(tmp_path, make_pdf):
    """Un escaneo sin capa de texto se reporta, no se indexa vacío."""
    make_pdf(tmp_path / "escaneo.pdf", [""])

    with pytest.raises(ExtractionError, match="OCR"):
        extract_pdf(tmp_path / "escaneo.pdf")


def test_corrupt_pdf_raises_extraction_error(tmp_path):
    (tmp_path / "roto.pdf").write_bytes(b"%PDF-1.4 basura")

    with pytest.raises(ExtractionError):
        extract_pdf(tmp_path / "roto.pdf")


def test_pdf_line_breaks_are_rejoined_into_paragraphs(tmp_path, make_pdf):
    """El PDF corta líneas por layout; unirlas evita falsos límites de párrafo."""
    make_pdf(tmp_path / "doc.pdf", ["Una frase que el layout parte a la mitad"])

    text = extract_pdf(tmp_path / "doc.pdf")
    body = text.split("\n\n", 1)[1]

    assert "\n" not in body.strip()


# --- DOCX --------------------------------------------------------------------


def test_docx_text_is_extracted(tmp_path, make_docx):
    make_docx(tmp_path / "doc.docx", [(None, "El mantenimiento es trimestral.")])

    text = extract_docx(tmp_path / "doc.docx")

    assert "El mantenimiento es trimestral." in text


def test_docx_heading_styles_become_markdown_headers(tmp_path, make_docx):
    make_docx(
        tmp_path / "doc.docx",
        [
            ("Heading 1", "Manual"),
            ("Heading 2", "Mantenimiento"),
            (None, "Revisar las hélices."),
        ],
    )

    text = extract_docx(tmp_path / "doc.docx")

    assert "# Manual" in text
    assert "## Mantenimiento" in text


def test_docx_hierarchy_survives_into_citations(tmp_path, make_docx):
    """Los estilos de Word producen la misma jerarquía que un Markdown nativo."""
    body = "Las hélices se inspeccionan buscando fisuras antes de cada vuelo. " * 3
    make_docx(
        tmp_path / "manual.docx",
        [("Heading 1", "Manual"), ("Heading 2", "Mantenimiento"), (None, body)],
    )

    document = Document(source_file="manual.docx", content=extract_docx(tmp_path / "manual.docx"))
    chunks = chunk_document(document, ChunkingConfig())

    assert any(c.citation == "manual.docx § Manual > Mantenimiento" for c in chunks)


def test_docx_empty_paragraphs_are_dropped(tmp_path, make_docx):
    make_docx(tmp_path / "doc.docx", [(None, "Contenido."), (None, "   "), (None, "Más.")])

    text = extract_docx(tmp_path / "doc.docx")

    assert "\n\n\n" not in text


def test_docx_tables_are_rendered_as_markdown(tmp_path):
    """Una tabla no se descarta: su contenido tiene que quedar buscable."""
    import docx

    document = docx.Document()
    document.add_paragraph("Antes de la tabla.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Componente"
    table.cell(0, 1).text = "Intervalo"
    table.cell(1, 0).text = "Batería"
    table.cell(1, 1).text = "50 ciclos"
    document.save(str(tmp_path / "doc.docx"))

    text = extract_docx(tmp_path / "doc.docx")

    assert "| Componente | Intervalo |" in text
    assert "| Batería | 50 ciclos |" in text


def test_docx_preserves_order_of_paragraphs_and_tables(tmp_path):
    """El intercalado importa: una tabla bajo su encabezado hereda esa sección."""
    import docx

    document = docx.Document()
    document.add_paragraph("Sección A", style="Heading 1")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Dato de A"
    document.add_paragraph("Sección B", style="Heading 1")
    document.save(str(tmp_path / "doc.docx"))

    text = extract_docx(tmp_path / "doc.docx")

    assert text.index("# Sección A") < text.index("Dato de A") < text.index("# Sección B")


def test_corrupt_docx_raises_extraction_error(tmp_path):
    (tmp_path / "roto.docx").write_bytes(b"PK\x03\x04 no es un docx")

    with pytest.raises(ExtractionError):
        extract_docx(tmp_path / "roto.docx")


# --- Registro ----------------------------------------------------------------


def test_registry_covers_the_documented_formats():
    assert set(EXTRACTORS) == {".md", ".markdown", ".txt", ".pdf", ".docx"}
